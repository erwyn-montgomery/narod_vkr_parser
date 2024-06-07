import pandas as pd
import olefile # для WordPadawan и ExcelPadawan
import docx # для WordPadawan
import openpyxl # для ExcelPadawan
import pptx # для PPtPadawan
from helper_functions.decode_str import decode_str
import pathlib


class WordPadawan:

    def __init__(self):
        self._common = False
        self._ex = ["doc", "docx"]
        self._doc = None
        self._doc_path = None
        self._error_log = {}

    def load_doc(self, path):
        try:
            self._doc = docx.Document(path)
            self._doc_path = path
        except Exception as e:
            with open("log_file.txt", "a") as f:
                f.write(f"Error | word | {path} | {e}\n")

    @property
    def doc(self):
        return self._doc
    
    @property
    def error_log(self):
        return self._error_log
    
    def get_doc_text(self):
        try:
            text = ""
            for par in self.doc.paragraphs:
                text += "\n" + par.text
            text = f"'{text}'"
            return text
        except Exception as e:
            pass
    
    def get_meta_inf(self, path):
        meta = {}
        ext = pathlib.Path(path).suffix[1:].lower()
        if ext == "doc":
            try:
                with olefile.OleFileIO(path) as ole:
                    meta_data = ole.get_metadata()
                    if meta_data.title:
                        meta["title"] = decode_str(meta_data.title)
                    if meta_data.author:
                        meta["author"] = decode_str(meta_data.author)
                    if meta_data.num_pages:
                        meta["page_count"] = meta_data.num_pages
                    if meta_data.num_chars:
                        meta["char_count"] = meta_data.num_chars
                    if meta_data.num_words:
                        meta["word_count"] = meta_data.num_words
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | word | {path} | {e}\n")
        if ext == "docx":
            try:
                self.load_doc(path)
                meta_data = self.doc.core_properties
                if meta_data.title:
                    meta["title"] = meta_data.title
                if meta_data.author:
                    meta["author"] = meta_data.author
                if self.get_doc_text():
                    meta["text"] = self.get_doc_text()
                    meta["char_count"] = len(meta["text"])
                    meta["word_count"] = len(meta["text"].split())
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | word | {path} | {e}\n")
        return meta
    

class ExcelPadawan:

    def __init__(self):
        self._common = False
        self._ex = ["xls", "xlsx"]
        self._df = None
        self._error_log = {}

    def load_df(self, path):
        try:
            self._df = pd.read_excel(path)
        except Exception as e:
            pass 

    @property
    def df(self):
        return self._df
    
    @property
    def error_log(self):
        return self._error_log    

    def get_meta_inf(self, path):
        meta = {}
        self.load_df(path)
        ext = pathlib.Path(path).suffix[1:].lower()
        if ext == "xls":
            try:
                with olefile.OleFileIO(path) as ole:
                    meta_data = ole.get_metadata()
                    if meta_data.title:
                        meta["title"] = decode_str(meta_data.title)
                    if meta_data.author:
                        meta["author"] = decode_str(meta_data.author)
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | excel | {path} | {e}\n")
        if ext == "xlsx":
            try:
                file = openpyxl.load_workbook(path)
                if file.properties.title:
                    meta["title"] = file.properties.title
                if file.properties.creator:
                    meta["author"] = file.properties.creator
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | excel | {path} | {e}\n")
        try:
            meta["rows"] = self.df.shape[0]
            meta["columns"] = self.df.shape[1]
        except Exception as e:
            with open("log_file.txt", "a") as f:
                    f.write(f"Error | excel | {path} | {e}\n")
        return meta
    

class PPtPadawan:

    def __init__(self):
        self._common = False
        self._ex = ["ppt", "pptx"]
        self._doc = None
        self._error_log = {}

    def load_doc(self, path):
        try:
            self._doc = pptx.Presentation(path)
        except Exception as e:
            with open("log_file.txt", "a") as f:
                f.write(f"Error | ppt | {path} | {e}\n")
    
    @property
    def doc(self):
        return self._doc
    
    @property
    def error_log(self):
        return self._error_log
    
    def get_meta_inf(self, path):
        meta = {}
        ext = pathlib.Path(path).suffix[1:].lower()
        if ext == "ppt":
            try:
                with olefile.OleFileIO(path) as ole:
                    meta_data = ole.get_metadata()
                    if meta_data.title:
                        meta["title"] = decode_str(meta_data.title)
                    if meta_data.author:
                        meta["author"] = decode_str(meta_data.author)
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | ppt | {path} | {e}\n")
        if ext == "pptx":
            try:
                self.load_doc(path)
                meta_data = self.doc.core_properties
                if meta_data.title:
                    meta["title"] = meta_data.title
                if meta_data.author:
                    meta["author"] = meta_data.author
                try:
                    meta["slides_count"] = len(self.doc.slides)
                except Exception as e:
                    with open("log_file.txt", "a") as f:
                        f.write(f"Error | ppt | {path} | {e}\n")
            except Exception as e:
                with open("log_file.txt", "a") as f:
                    f.write(f"Error | ppt | {path} | {e}\n")
        return meta